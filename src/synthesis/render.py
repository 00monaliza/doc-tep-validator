"""Render `Document` to a text-layer PDF (fpdf2 + DejaVu Sans) and to DOCX.

fpdf2 was chosen over reportlab after scripts/check_env.py showed both render
all Kazakh letters correctly with an embedded TTF; fpdf2 has a native table API.
"""

from __future__ import annotations

import re
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fpdf import FPDF, FontFace

from src.synthesis.document import Block, Document

ROOT = Path(__file__).resolve().parents[2]
FONT_REGULAR = ROOT / "assets" / "fonts" / "DejaVuSans.ttf"
FONT_BOLD = ROOT / "assets" / "fonts" / "DejaVuSans-Bold.ttf"

NUMERIC = re.compile(r"^-?[\d ]+(,\d+)?$")
SHEET = {"ru": "Лист", "kz": "Парақ"}


def _numeric_columns(block: Block) -> list[bool]:
    cols = len(block.header)
    return [
        all(NUMERIC.match(r[c]) for r in block.rows if r[c]) and any(r[c] for r in block.rows)
        for c in range(cols)
    ]


class _PDF(FPDF):
    def __init__(self, doc: Document):
        super().__init__(format="A4")
        self.doc = doc
        self.add_font("DejaVu", "", str(FONT_REGULAR))
        self.add_font("DejaVu", "B", str(FONT_BOLD))
        self.set_margins(20, 18, 12)
        self.set_auto_page_break(True, margin=18)

    def header(self) -> None:
        self.set_font("DejaVu", size=7)
        self.cell(0, 4, f"{self.doc.code}    {self.doc.title}", align="R", new_x="LMARGIN", new_y="NEXT")
        self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
        self.ln(3)

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("DejaVu", size=7)
        self.cell(0, 4, f"{SHEET[self.doc.lang]} {self.page_no()}", align="R")


def render_pdf(doc: Document, path: Path) -> None:
    pdf = _PDF(doc)
    pdf.add_page()
    width = pdf.epw
    for block in doc.blocks:
        if block.kind == "title":
            pdf.set_font("DejaVu", "B", 13)
            pdf.multi_cell(width, 6.5, doc.title, align="C", new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("DejaVu", "", 10.5)
            pdf.multi_cell(width, 5.5, block.text, align="C", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(4)
        elif block.kind == "heading":
            pdf.ln(2)
            pdf.set_font("DejaVu", "B", 11)
            pdf.multi_cell(width, 6, block.text, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif block.kind == "para":
            pdf.set_font("DejaVu", "", 10)
            pdf.multi_cell(width, 5, block.text, align="J", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1.5)
        elif block.kind == "table":
            pdf.set_font("DejaVu", "", 8.5)
            numeric = _numeric_columns(block)
            align = tuple("RIGHT" if n else "LEFT" for n in numeric)
            bold = FontFace(emphasis="BOLD")
            with pdf.table(col_widths=tuple(block.col_widths), text_align=align, line_height=4.4,
                           headings_style=FontFace(emphasis="BOLD", fill_color=(235, 235, 235)),
                           width=width) as table:
                head = table.row()
                for h in block.header:
                    head.cell(h, align="CENTER")
                for i, row in enumerate(block.rows):
                    r = table.row()
                    for cell in row:
                        r.cell(cell, style=bold if i in block.bold_rows else None)
            pdf.ln(3)
    pdf.output(str(path))


def render_docx(doc: Document, path: Path) -> None:
    d = docx.Document()
    normal = d.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    d.sections[0].header.paragraphs[0].text = f"{doc.code}    {doc.title}"
    for block in doc.blocks:
        if block.kind == "title":
            for text, size in ((doc.title, 14), (block.text, 11)):
                p = d.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(text)
                run.bold = size == 14
                run.font.size = Pt(size)
        elif block.kind == "heading":
            d.add_heading(block.text, level=2)
        elif block.kind == "para":
            d.add_paragraph(block.text)
        elif block.kind == "table":
            table = d.add_table(rows=1 + len(block.rows), cols=len(block.header))
            table.style = "Table Grid"
            numeric = _numeric_columns(block)
            for c, h in enumerate(block.header):
                cell = table.rows[0].cells[c]
                cell.text = h
                cell.paragraphs[0].runs[0].bold = True
            for i, row in enumerate(block.rows, start=1):
                for c, text in enumerate(row):
                    cell = table.rows[i].cells[c]
                    cell.text = text
                    par = cell.paragraphs[0]
                    if numeric[c]:
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    if (i - 1) in block.bold_rows and par.runs:
                        par.runs[0].bold = True
            d.add_paragraph()
    d.save(str(path))
