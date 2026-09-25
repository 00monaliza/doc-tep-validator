"""Paragraph and table extraction from DOCX (python-docx)."""

from __future__ import annotations

from pathlib import Path

import docx


def extract_paragraphs(path: str | Path) -> list[str]:
    return [p.text for p in docx.Document(str(path)).paragraphs if p.text.strip()]


def extract_tables(path: str | Path) -> list[list[list[str]]]:
    document = docx.Document(str(path))
    return [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]


def extract_text(path: str | Path) -> str:
    parts = extract_paragraphs(path)
    for table in extract_tables(path):
        parts.extend(" | ".join(row) for row in table)
    return "\n".join(parts)
